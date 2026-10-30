"""Tests for the parsing layer. They read local files only; no network.

The most important group comes first: the Markdown / docx / PDF versions of the same document
parse into the same structure. The docx and PDF files live in tests/fixtures/documents/ and are
generated from the Markdown by scripts/make_sample_documents.py.
"""
from io import BytesIO
from pathlib import Path

import docx
import pytest
from docx.enum.style import WD_STYLE_TYPE

from app.rag.parsing import (
    DocumentParseError,
    ElementKind,
    UnsupportedFormatError,
    parse_file,
    parser_for,
)
from app.rag.parsing.docx_parser import DocxParser
from app.rag.parsing.markdown_parser import MarkdownParser
from app.rag.parsing.pdf_parser import PdfParser, _Assembler, _Line

REPO_ROOT = Path(__file__).resolve().parents[2]
FIXTURES = Path(__file__).resolve().parent / "fixtures" / "documents"

SAMPLES = [
    REPO_ROOT / "sop_docs" / "account-restriction-handling.md",            # two list levels
    REPO_ROOT / "sop_docs" / "host-onboarding-and-review-electronics.md",
    REPO_ROOT / "campaign_docs" / "beauty-black-friday-live-review.md",     # tables, one split across pages
]
REVIEW_PDF = FIXTURES / "beauty-black-friday-live-review.pdf"


def _structure(document) -> list[tuple]:
    """What gets compared: kind, level, numbered or not, text, heading path, table cells.

    Text is compared exactly, spaces included. A PDF contains no space characters; spaces are
    inferred from the gaps between letters, and a wrong inference glues words together.
    Nested list items are not compared on "numbered or not": when the fixtures were generated,
    WPS turned the bullet list nested inside a numbered list into a numbered one. That is the
    converter changing the content, not a parser problem.
    """
    return [
        (e.kind, e.level,
         e.ordered if e.kind is not ElementKind.LIST_ITEM or e.level == 1 else None,
         e.text, e.heading_path, e.rows)
        for e in document.elements
    ]


# ---------------------------------------------------------------------------
# All three formats agree
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("suffix", [".docx", ".pdf"])
@pytest.mark.parametrize("markdown", SAMPLES, ids=lambda path: path.stem)
def test_same_structure_as_markdown(markdown: Path, suffix: str):
    expected = parse_file(markdown)
    actual = parse_file(FIXTURES / f"{markdown.stem}{suffix}")
    assert _structure(actual) == _structure(expected)


def test_fixtures_exist_for_every_sample():
    # Without a fixture the parametrized test above fails with FileNotFoundError; this says what to do
    missing = [f"{md.stem}{suffix}" for md in SAMPLES for suffix in (".docx", ".pdf")
               if not (FIXTURES / f"{md.stem}{suffix}").is_file()]
    assert not missing, f"missing fixtures, regenerate them with scripts/make_sample_documents.py: {missing}"


def test_heading_levels_and_paths():
    document = parse_file(SAMPLES[0])
    title = "Account Restriction Handling"
    headings = [(e.level, e.text, e.heading_path) for e in document.elements if e.kind is ElementKind.HEADING]
    assert headings == [
        (1, title, ()),
        (2, "When to Use", (title,)),
        (2, "Procedure", (title,)),
        (2, "Partner Teams", (title,)),
    ]
    steps = [e for e in document.elements if e.heading_path == (title, "Procedure")]
    assert [(e.level, e.ordered) for e in steps] == [(1, True)] * 7 + [(2, False)] * 2
    assert steps[0].text.startswith("Identify the rule and the account")   # the "1." is not part of the text


# ---------------------------------------------------------------------------
# What the PDF parser has to infer
# ---------------------------------------------------------------------------
def test_pdf_page_numbers():
    pdf = parse_file(REVIEW_PDF)
    assert all(e.page is not None for e in pdf.elements)
    findings = next(e for e in pdf.elements if e.text == "Findings:")
    assert findings.page == 2
    assert all(e.page is None for e in parse_file(FIXTURES / "beauty-black-friday-live-review.docx").elements)


def test_pdf_table_split_across_pages_is_merged():
    # Only the header row of the product table fits on page 1; WPS repeats it on page 2 above the data rows
    tables = [e for e in parse_file(REVIEW_PDF).elements if e.kind is ElementKind.TABLE]
    products = tables[1]
    assert products.page == 1   # the table starts on page 1
    assert [row[0] for row in products.rows] == ["Product", "Sheet Mask", "Lipstick", "Serum"]


def test_pdf_words_on_tightly_set_lines_stay_separate():
    # On this line the gap between words is only 2.6pt; pdfplumber's default tolerance (3pt)
    # read it as "byhandorimprovisesaformula"
    pdf = parse_file(REVIEW_PDF)
    assert any("by hand or improvises a formula." in e.text for e in pdf.elements)


def test_pdf_inline_code_stays_in_its_sentence():
    # Inline code uses a smaller font and sits lower; split by top coordinate it would become a
    # line of its own somewhere else. The spacing around it must not leak into the text either
    pdf = parse_file(REVIEW_PDF)
    paragraph = next(e for e in pdf.elements if e.text.startswith("The first two versions were wrong."))
    assert "product cost is already deducted in campaign_products.profit (profit = sales - product cost;" \
        in paragraph.text
    assert any(e.text.startswith("Actual operating cost (campaigns.actual_operating_cost): ") for e in pdf.elements)


def test_pdf_font_size_noise_does_not_create_heading_levels():
    # The fixtures contain sizes like 18.0, 17.999999999999986 and 18.00000000000003 for the same
    # headings. Unrounded, each would become its own heading level
    pdf = parse_file(REVIEW_PDF)
    assert sorted({e.level for e in pdf.elements if e.kind is ElementKind.HEADING}) == [1, 2]


def test_pdf_bullet_glyph_is_not_kept_in_text():
    # Word / WPS draw bullets with a private-use character from the Symbol font; left in the text it is garbage
    for markdown in SAMPLES:
        pdf = parse_file(FIXTURES / f"{markdown.stem}.pdf")
        assert not any("" in e.text for e in pdf.elements)


@pytest.mark.parametrize("last_line_on_page_1, second_page, expected", [
    # The last line of page 1 does not end a sentence: the paragraph continues on page 2
    ("Appeal evidence includes the stream recording, product certificates", "and shipping records.",
     ["Appeal evidence includes the stream recording, product certificates and shipping records."]),
    # It ends with a period: page 2 starts a new paragraph
    ("The first paragraph ends here.", "The second paragraph starts on a new page.",
     ["The first paragraph ends here.", "The second paragraph starts on a new page."]),
])
def test_pdf_paragraph_across_page_break(last_line_on_page_1, second_page, expected):
    # No fixture has a paragraph across a page break, so the rule is tested on lines directly
    page_1 = [_Line(text=last_line_on_page_1, x0=90, top=760, bottom=772, size=12.0, page=1)]
    page_2 = [_Line(text=second_page, x0=90, top=72, bottom=84, size=12.0, page=2)]
    elements = _Assembler(page_1 + page_2).run([(page_1, []), (page_2, [])])
    assert [e.text for e in elements] == expected


def test_pdf_without_text_layer_is_rejected():
    # A blank page with no text at all, which is what a scan looks like (a scan holds only images)
    blank = (b"%PDF-1.4\n1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n"
             b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n"
             b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] >> endobj\n"
             b"trailer << /Root 1 0 R >>\n%%EOF\n")
    with pytest.raises(DocumentParseError, match="no text layer"):
        PdfParser().parse(blank)


def test_broken_pdf_is_rejected():
    with pytest.raises(DocumentParseError, match="cannot read the PDF"):
        PdfParser().parse(b"this is not a pdf")


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------
def _md(text: str):
    return MarkdownParser().parse(text.encode("utf-8"))


def test_markdown_frontmatter_becomes_metadata_not_text():
    document = parse_file(SAMPLES[0])
    assert document.metadata["department"] == "ALL"
    assert not any("responsible_team" in e.text for e in document.elements)


def test_markdown_heading_levels_are_relative():
    # Only ## and ### are used; they become 1 and 2, matching a PDF where only "large / smaller font" is visible
    document = _md("## General\n\nBody one\n\n### Details\n\nBody two\n\n## Appendix\n\nBody three\n")
    assert [(e.kind, e.level, e.heading_path) for e in document.elements] == [
        (ElementKind.HEADING, 1, ()),
        (ElementKind.PARAGRAPH, 0, ("General",)),
        (ElementKind.HEADING, 2, ("General",)),
        (ElementKind.PARAGRAPH, 0, ("General", "Details")),
        (ElementKind.HEADING, 1, ()),                  # back to level 1, leaving "General > Details"
        (ElementKind.PARAGRAPH, 0, ("Appendix",)),
    ]


def test_markdown_line_breaks_inside_a_paragraph():
    document = _md("a paragraph wrapped\nby hand\n\nsecond paragraph\n")
    assert [e.text for e in document.elements] == ["a paragraph wrapped by hand", "second paragraph"]


def test_markdown_table_keeps_empty_cells():
    document = _md("| Metric | Note |\n|---|---|\n| GMV | |\n")
    assert document.elements[0].rows == (("Metric", "Note"), ("GMV", ""))


def test_markdown_must_be_utf8():
    # Windows-1252, the old default on English Windows: "é" is the single byte 0xE9, invalid in UTF-8
    with pytest.raises(DocumentParseError, match="UTF-8"):
        MarkdownParser().parse("# Café opening checklist".encode("cp1252"))


def test_markdown_bad_frontmatter_is_rejected():
    with pytest.raises(DocumentParseError, match="YAML"):
        _md("---\ntitle: [unclosed\n---\nBody\n")


def test_empty_document_is_rejected():
    with pytest.raises(DocumentParseError, match="no content"):
        _md("---\ntitle: frontmatter only\n---\n")


# ---------------------------------------------------------------------------
# docx (built on the spot with python-docx, no fixture files)
# ---------------------------------------------------------------------------
def _docx(build) -> bytes:
    document = docx.Document()
    build(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_title_and_heading_styles():
    def build(d):
        d.add_heading("Employee Handbook", level=0)   # the "Title" style, one above "Heading 1"
        d.add_heading("Attendance", level=1)
        d.add_paragraph("Clock in by 9 a.m.")
    document = DocxParser().parse(_docx(build))
    assert [(e.kind, e.level, e.heading_path) for e in document.elements] == [
        (ElementKind.HEADING, 1, ()),
        (ElementKind.HEADING, 2, ("Employee Handbook",)),
        (ElementKind.PARAGRAPH, 0, ("Employee Handbook", "Attendance")),
    ]


def test_docx_custom_style_based_on_heading():
    def build(d):
        style = d.styles.add_style("Company Heading", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = d.styles["Heading 1"]
        d.add_paragraph("A heading with a custom style", style="Company Heading")
        d.add_paragraph("Body")
    document = DocxParser().parse(_docx(build))
    assert document.elements[0].kind is ElementKind.HEADING


def test_docx_numbering_from_list_styles():
    # The built-in "List Number / List Bullet" styles carry the numbering on the style, not the paragraph
    def build(d):
        d.add_paragraph("Step one", style="List Number")
        d.add_paragraph("A note", style="List Bullet")
    document = DocxParser().parse(_docx(build))
    assert [(e.kind, e.level, e.ordered, e.text) for e in document.elements] == [
        (ElementKind.LIST_ITEM, 1, True, "Step one"),
        (ElementKind.LIST_ITEM, 1, False, "A note"),
    ]


def test_docx_table_in_document_order():
    def build(d):
        d.add_paragraph("Before the table")
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Product", "Units sold"
        table.cell(1, 0).text, table.cell(1, 1).text = "Sheet Mask", "470"
        d.add_paragraph("After the table")
    document = DocxParser().parse(_docx(build))
    assert [e.kind for e in document.elements] == [ElementKind.PARAGRAPH, ElementKind.TABLE, ElementKind.PARAGRAPH]
    assert document.elements[1].rows == (("Product", "Units sold"), ("Sheet Mask", "470"))


def test_broken_docx_is_rejected():
    with pytest.raises(DocumentParseError, match="cannot open the docx"):
        DocxParser().parse(b"PK\x03\x04 not really a zip")


# ---------------------------------------------------------------------------
# Choosing the parser by extension
# ---------------------------------------------------------------------------
def test_parser_chosen_by_extension_case_insensitive():
    # Files in the upload folder are named "<uploadId>-<original name>"
    assert isinstance(parser_for("0f8c-account-restriction-handling.PDF"), PdfParser)
    assert isinstance(parser_for("a.Docx"), DocxParser)
    assert isinstance(parser_for("a.md"), MarkdownParser)


@pytest.mark.parametrize("filename", ["legacy-document.doc", "notes.txt", "no-extension"])
def test_unsupported_formats_are_rejected(filename: str):
    with pytest.raises(UnsupportedFormatError, match="docx"):
        parser_for(filename)
